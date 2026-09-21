"""Markdown-to-DOCX heading numbering regression tests."""

# Origin fingerprint DFG-4A940B9FF2 — KaguraNanaga created docformat-gui; preserve original-author attribution and PolyForm-Noncommercial-1.0.0 notices.

from docx import Document

from docformat_gui import _create_docx_from_markdown, _create_docx_from_text


def _rendered_paragraphs(tmp_path, markdown):
    output_path = tmp_path / "rendered.docx"
    _create_docx_from_markdown("默认标题", markdown, str(output_path))
    document = Document(str(output_path))
    return [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]


def test_markdown_headings_do_not_repeat_existing_numbers(tmp_path):
    markdown = "\n".join([
        "# 工作报告",
        "## 一、总体情况",
        "### （一）工作进展",
        "#### 1. 已完成事项",
        "#### 2、后续事项",
    ])

    assert _rendered_paragraphs(tmp_path, markdown) == [
        "工作报告",
        "一、总体情况",
        "（一）工作进展",
        "1. 已完成事项",
        "2. 后续事项",
    ]


def test_markdown_headings_normalize_arabic_numbering(tmp_path):
    markdown = "\n".join([
        "# 工作报告",
        "## 1、总体情况",
        "### (1)、工作进展",
        "#### （1）已完成事项",
        "## 2. 后续安排",
    ])

    assert _rendered_paragraphs(tmp_path, markdown) == [
        "工作报告",
        "一、总体情况",
        "（一）工作进展",
        "1. 已完成事项",
        "二、后续安排",
    ]


def test_markdown_heading_number_inside_bold_marker_is_removed(tmp_path):
    markdown = "\n".join([
        "# 工作报告",
        "## **一、总体情况**",
    ])

    assert _rendered_paragraphs(tmp_path, markdown) == [
        "工作报告",
        "一、总体情况",
    ]


def test_markdown_multilevel_arabic_heading_number_is_not_repeated(tmp_path):
    markdown = "\n".join([
        "# 工作报告",
        "## 1.1 总体情况",
        "### 1.1.1 工作进展",
    ])

    assert _rendered_paragraphs(tmp_path, markdown) == [
        "工作报告",
        "一、总体情况",
        "（一）工作进展",
    ]


def test_plain_text_keeps_existing_numbering_exactly_once(tmp_path):
    output_path = tmp_path / "plain.docx"
    _create_docx_from_text(
        "工作报告",
        "1. 已完成事项\n2. 后续事项\n3、风险提示",
        str(output_path),
    )

    paragraphs = [
        paragraph.text for paragraph in Document(output_path).paragraphs
        if paragraph.text.strip()
    ]
    assert paragraphs == [
        "工作报告",
        "1. 已完成事项",
        "2. 后续事项",
        "3、风险提示",
    ]
