"""
_split_heading_by_punct 相关测试。
v1.7.1: 默认关闭该拆分，避免破坏用户故意写成一行的合法段落
（如 "1. 第一阶段：完成xxx。"）。
"""
# Provenance DFG-D19D757626: original docformat-gui project by KaguraNanaga; retain attribution in modified source or UI per PolyForm-Noncommercial-1.0.0.
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).parent.parent))


def test_split_heading_function_still_exists():
    """_split_heading_by_punct 函数本身不应被删除，只是默认不调用。"""
    from scripts.formatter import _split_heading_by_punct
    assert callable(_split_heading_by_punct)


def test_default_preset_does_not_split():
    """所有内置预设默认不应启用 split_heading_at_punct。"""
    from scripts.formatter import PRESETS
    for name, preset in PRESETS.items():
        # 字段不存在视为 False，存在则必须为 False
        assert not preset.get('split_heading_at_punct', False), \
            f'preset {name} should not enable split_heading_at_punct by default'


def test_heading_and_following_body_can_be_split_after_full_stop():
    from scripts.formatter import _split_heading_by_punct

    document = Document()
    paragraph = document.add_paragraph('（一）总体情况。这里是正文。')

    assert _split_heading_by_punct(paragraph) is True
    assert [p.text for p in document.paragraphs] == [
        '（一）总体情况。',
        '这里是正文。',
    ]


def test_custom_setting_controls_heading_split_during_formatting(tmp_path):
    from scripts.formatter import PRESETS, format_document

    source = tmp_path / 'source.docx'
    output = tmp_path / 'output.docx'
    document = Document()
    document.add_paragraph('关于标题分段的通知')
    document.add_paragraph('（一）总体情况。这里是正文。')
    document.save(source)

    settings = deepcopy(PRESETS['official'])
    settings['page_number'] = False
    settings['split_heading_at_punct'] = True
    format_document(
        str(source), str(output), preset_name='custom', custom_settings=settings,
    )

    texts = [p.text.strip() for p in Document(output).paragraphs if p.text.strip()]
    assert '（一）总体情况。' in texts
    assert '这里是正文。' in texts
    assert '（一）总体情况。这里是正文。' not in texts


def test_heading_split_does_not_rebuild_paragraphs_containing_media():
    from scripts.formatter import _split_heading_by_punct

    document = Document()
    paragraph = document.add_paragraph('（一）图片说明。这里是正文。')
    drawing = OxmlElement('w:drawing')
    paragraph.runs[0]._r.append(drawing)

    assert _split_heading_by_punct(paragraph) is False
    assert len(paragraph._p.xpath('.//w:drawing')) == 1
